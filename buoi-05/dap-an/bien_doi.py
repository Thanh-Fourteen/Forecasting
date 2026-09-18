# %% [markdown]
# # Buổi 5 — Điều chỉnh lịch, dân số, lạm phát; Box-Cox và hiệu chỉnh bias khi đổi ngược
#
# Bản ĐÃ SỬA. Dữ liệu: doanh số bán lẻ và dịch vụ ăn uống Mỹ (Census MARTS), CPI-U (BLS), dân số (BEA NIPA).

# %%
from __future__ import annotations

import re

import numpy as np
import openpyxl
import pandas as pd

import tv

TONG = "Retail and food services sales, total"


def doc_ban_le(ten: str = TONG, da_dieu_chinh: bool = False) -> pd.Series:
    """Doanh số theo tháng (triệu USD). `da_dieu_chinh=False` = khối NOT ADJUSTED; True = ADJUSTED(2)."""
    tep = tv.THU_MUC_DU_LIEU / "census-marts-ban-le" / "mrtssales92-present.xlsx"
    can_khoi = "ADJUSTED(2)" if da_dieu_chinh else "NOT ADJUSTED"
    gia_tri: dict[pd.Timestamp, float] = {}
    wb = openpyxl.load_workbook(tep, read_only=True, data_only=True)
    for sheet in wb.sheetnames:
        dong = list(wb[sheet].iter_rows(values_only=True))
        cot_thang = {}
        for j, nhan in enumerate(dong[4]):
            m = re.match(r"^([A-Z][a-z]{2})\.? (\d{4})", nhan) if isinstance(nhan, str) else None
            if m:  # "Jan. 2025", "May 2025", "Dec. 2025(p)"
                cot_thang[j] = pd.Timestamp(f"{m.group(2)}-{pd.to_datetime(m.group(1), format='%b').month:02d}-01")
        khoi = None
        for r in dong[5:]:
            if r[1] in ("NOT ADJUSTED", "ADJUSTED(2)"):
                khoi = r[1]
            elif khoi == can_khoi and isinstance(r[1], str) and r[1].strip() == ten:
                for j, thang in cot_thang.items():
                    gia_tri[thang] = pd.to_numeric(r[j], errors="coerce")  # "(S)", "(NA)" → NaN
    wb.close()
    return pd.Series(gia_tri, name=ten).sort_index().asfreq("MS")


def doc_cpi(ma: str = "CUUR0000SA0") -> pd.Series:
    """CPI-U theo tháng (1982–84 = 100). CUUR0000SA0 = chưa điều chỉnh mùa vụ; bỏ M13 (trung bình năm)."""
    df = pd.read_csv(tv.THU_MUC_DU_LIEU / "bls-cpi-u" / "cu.data.1.AllItems.tsv", sep="\t", dtype=str)
    df.columns = df.columns.str.strip()
    df = df[(df["series_id"].str.strip() == ma) & (df["period"] != "M13")]
    thang = pd.to_datetime(df["year"].str.strip() + "-" + df["period"].str[1:] + "-01")
    gia_tri = pd.to_numeric(df["value"].str.strip(), errors="coerce")  # "-" = tháng không công bố
    return pd.Series(gia_tri.to_numpy(), index=thang, name=ma).sort_index().asfreq("MS")


def doc_dan_so() -> pd.Series:
    """Dân số Mỹ giữa kỳ theo tháng (nghìn người), BEA NIPA mã B230RC."""
    df = pd.read_csv(tv.THU_MUC_DU_LIEU / "bea-nipa-thang" / "NipaDataM.txt", thousands=",")
    df = df[df["%SeriesCode"] == "B230RC"]
    thang = pd.to_datetime(df["Period"].str.replace("M", "-") + "-01")
    return pd.Series(df["Value"].astype(float).to_numpy(), index=thang, name="dan_so").sort_index().asfreq("MS")


# %% [markdown]
# ## Điều chỉnh

# %%
def theo_ngay(y: pd.Series) -> pd.Series:
    """Điều chỉnh lịch: chia tổng tháng cho số ngày của tháng."""
    return y / y.index.days_in_month


def so_sanh_thang(y: pd.Series, thang_truoc: str, thang_sau: str) -> float:
    """Tăng trưởng (tỷ lệ) từ tháng trước sang tháng sau, SAU khi điều chỉnh số ngày."""
    ngay = theo_ngay(y)
    return float(ngay[pd.Timestamp(thang_sau)] / ngay[pd.Timestamp(thang_truoc)] - 1)


def gia_thuc(y: pd.Series, cpi: pd.Series, nam_goc: int) -> pd.Series:
    """Đổi sang giá của năm gốc: x_t = y_t / CPI_t × CPI_trung_bình(năm gốc)."""
    cpi = cpi.interpolate(limit=1, limit_area="inside")  # BLS không công bố 10/2025 → nội suy MỘT tháng, ghi rõ
    return y / cpi.reindex(y.index) * cpi[str(nam_goc)].mean()


def tang_truong(y: pd.Series, tu_nam: int, den_nam: int) -> float:
    """Tỷ lệ tăng của tổng năm `den_nam` so với tổng năm `tu_nam`."""
    return float(y[str(den_nam)].sum() / y[str(tu_nam)].sum() - 1)


def tang_truong_thuc_dau_nguoi(y: pd.Series, cpi: pd.Series, dan_so: pd.Series, tu_nam: int, den_nam: int) -> float:
    """Tăng trưởng sau khi điều chỉnh CẢ lạm phát lẫn dân số."""
    thuc = gia_thuc(y, cpi, den_nam) / dan_so.reindex(y.index)
    return float(thuc[str(den_nam)].mean() / thuc[str(tu_nam)].mean() - 1)


# %% [markdown]
# ## Box-Cox tự viết

# %%
def boxcox(y, lam: float) -> np.ndarray:
    """w = log(y) nếu λ = 0; (sign(y)|y|^λ − 1)/λ nếu λ ≠ 0 (dạng của FPP, Bickel & Doksum 1981)."""
    y = np.asarray(y, dtype=float)
    if lam == 0:
        return np.log(y)
    return (np.sign(y) * np.abs(y) ** lam - 1) / lam


def boxcox_nguoc(w, lam: float, sigma2=0.0) -> np.ndarray:
    """Đổi ngược; sigma2 > 0 → hiệu chỉnh bias (FPP công thức 5.3) để ra TRUNG BÌNH thay vì trung vị."""
    w = np.asarray(w, dtype=float)
    sigma2 = np.asarray(sigma2, dtype=float)
    if lam == 0:
        return np.exp(w) * (1 + sigma2 / 2)
    goc = lam * w + 1
    return np.sign(goc) * np.abs(goc) ** (1 / lam) * (1 + sigma2 * (1 - lam) / (2 * goc**2))


def guerrero(y, m: int, can=(-1.0, 2.0), buoc: int = 3001) -> float:
    """λ Guerrero (1993): chia chuỗi thành các khối m quan sát liền (bỏ khối đầu thiếu), tính s_i/μ_i^(1−λ) mỗi khối,
    chọn λ làm hệ số biến thiên (sd/mean) của các tỷ số đó nhỏ nhất. Tìm trên lưới đều trong `can`."""
    y = np.asarray(y, dtype=float)
    so_khoi = y.size // m
    khoi = y[y.size - so_khoi * m:].reshape(so_khoi, m)
    tb, sd = khoi.mean(axis=1), khoi.std(axis=1, ddof=1)
    luoi = np.linspace(can[0], can[1], buoc)
    ty_so = sd[None, :] / tb[None, :] ** (1 - luoi[:, None])
    cv = ty_so.std(axis=1, ddof=1) / ty_so.mean(axis=1)
    return float(luoi[np.argmin(cv)])


# %% [markdown]
# ## Dự báo trên thang log, đổi ngược có/không hiệu chỉnh bias

# %%
def du_bao_log(y: pd.Series, goc: pd.Timestamp, so_buoc: int = 12, so_thang_hoc: int = 96) -> pd.DataFrame:
    """Seasonal naive có drift trên thang log: ŵ_{T+h} = w_{T+h−12} + drift, drift & σ² từ sai phân mùa vụ của cửa sổ học.
    Chỉ dùng dữ liệu TRƯỚC gốc. Trả cột trung_vi = exp(ŵ) và trung_binh = exp(ŵ)(1 + σ²/2)."""
    hoc = y[(y.index < goc) & (y.index >= goc - pd.DateOffset(months=so_thang_hoc))]
    w = np.log(hoc)
    sai_phan = (w - w.shift(12)).dropna()
    drift, s2 = sai_phan.mean(), sai_phan.var(ddof=1)
    thoi_gian = pd.date_range(goc, periods=so_buoc, freq="MS")
    w_hat = np.array([w[t - pd.DateOffset(years=1)] + drift for t in thoi_gian])
    return pd.DataFrame({"trung_vi": boxcox_nguoc(w_hat, 0.0), "trung_binh": boxcox_nguoc(w_hat, 0.0, s2),
                         "sigma2": s2}, index=thoi_gian)


def danh_gia_bias(y: pd.Series, goc_dau: str = "2012-01-01", goc_cuoi: str = "2018-12-01") -> dict[str, float]:
    """Gốc mỗi tháng, tầm 12: tổng dự báo / tổng thực tế − 1 (%) cho hai cách đổi ngược."""
    phan = []
    for goc in pd.date_range(goc_dau, goc_cuoi, freq="MS"):
        du_bao = du_bao_log(y, goc)
        phan.append(du_bao.assign(y=y.reindex(du_bao.index).to_numpy()))
    bang = pd.concat(phan)
    return {cot: float((bang[cot].sum() / bang["y"].sum() - 1) * 100) for cot in ("trung_vi", "trung_binh")}


# %%
if __name__ == "__main__":
    y = doc_ban_le()
    cpi, dan_so = doc_cpi(), doc_dan_so()
    print(f"bán lẻ {y.index[0]:%Y-%m} → {y.index[-1]:%Y-%m}, {len(y)} tháng")
    print("T3/T2 2023 thô:", round(float(y[pd.Timestamp("2023-03")] / y[pd.Timestamp("2023-02")] - 1), 4),
          "theo ngày:", round(so_sanh_thang(y, "2023-02", "2023-03"), 4))
    print("tăng 1993→2025 danh nghĩa:", round(tang_truong(y, 1993, 2025), 3),
          "thực/đầu người:", round(tang_truong_thuc_dau_nguoi(y, cpi, dan_so, 1993, 2025), 3))
    print("λ Guerrero (1992–2019, m=12):", guerrero(y[:"2019-12"], 12))
    print("bias % (tổng dự báo/tổng thực − 1):", {k: round(v, 2) for k, v in danh_gia_bias(y).items()})
